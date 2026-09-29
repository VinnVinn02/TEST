import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import normalize as N
from .assemble import COLUMNS
from .verify import verify_all

TEXT_COLS = {"NIK", "NO HP", "NO RANGKA", "NO MESIN"}


def bast_tag(bast_no: str) -> str:
    """047-FDB-2026-9-00024 -> 047-00024"""
    p = (bast_no or "").split("-")
    return f"{p[0]}-{p[-1]}" if len(p) >= 2 else (bast_no or "BAST")


def _sheet(wb, title, headers, rows, widths=None, wrap=()):
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        width = min(max([len(str(h))] + [max(len(x) for x in str(c.value or "").split("\n")) for c in ws[get_column_letter(i)][1:]]) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = width
        for c in ws[get_column_letter(i)][1:]:
            if h in TEXT_COLS:
                c.number_format = "@"
            if h in wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def _bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def refresh_filenames(records):
    for r in records:
        fn = N.file_names(r["NAMA LENGKAP"])
        r["NAMA FILE"] = "\n".join([fn["FAKTUR"], fn["KTP"], fn["SERTIFIKAT"]])
    return records


FILL = {"CEK": PatternFill("solid", fgColor="FFEB9C"), "KOSONG": PatternFill("solid", fgColor="F4B6B6")}


def data_berkas(records, warnings=(), general=(), cells=None):
    wb = Workbook()
    ws0 = _sheet(wb, "DATA KENDARAAN", COLUMNS, [[r[c] for c in COLUMNS] for r in records], wrap=("NAMA FILE",))
    for i, cf in enumerate(cells or [], 2):  # kuning = perlu dicek, merah = kosong
        for j, col in enumerate(COLUMNS, 1):
            if cf.get(col) in FILL:
                ws0.cell(i, j).fill = FILL[cf[col]]
    ws = wb.create_sheet("PERINGATAN")
    ws.append(["NAMA LENGKAP", "PERINGATAN"])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(["(LEGENDA)", "Sel KUNING di DATA KENDARAAN = dua OCR berbeda / tidak terkonfirmasi, WAJIB dicek dari gambar. "
                            "Sel MERAH = kosong / tidak ditemukan."])
    for g in general:
        ws.append(["(UMUM)", g])
    for r, ws_ in zip(records, warnings):
        for m in ws_:
            ws.append([r["NAMA LENGKAP"], m])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 100
    return _bytes(wb)


def rename_rows(records, meta, pdf_stem):
    prefix = pdf_stem.replace("_", " ")
    rows, n = [], 0
    for r, m in zip(records, meta):
        fn = N.file_names(r["NAMA LENGKAP"])
        for kind, label in (("FAKTUR", "FAKTUR KENDARAAN BERMOTOR"), ("SERTIFIKAT", "SERTIFIKAT NIK"), ("KTP", "KTP")):
            pg = m["pages"][kind]
            if pg:
                n += 1
                rows.append((n, pg, f"{prefix}_{pg}.jpg", fn[kind], label))
    return rows


def rename_xlsx(rows):
    wb = Workbook()
    _sheet(wb, "RENAME FILE", ["NO", "Nama Lama", "Nama Baru", "JENIS DOKUMEN"], [(a, c, d, e) for a, _, c, d, e in rows])
    return _bytes(wb)


REKAP_HEAD = ["NO", "NOMOR FAKTUR", "NAMA PEMILIK", "ALAMAT", "TYPE", "NOMOR RANGKA", "NOMOR MESIN", "TAHUN",
              "WARNA", "NO HP", "GMAIL"]


def rekap_rows(records, meta):
    rows = []
    for i, (r, m) in enumerate(zip(records, meta), 1):
        alamat = " - ".join(x for x in (r["ALAMAT 1"], r["ALAMAT 2"], "BATAM") if x)
        rows.append([i, r["NO. FAKTUR"], r["NAMA LENGKAP"], alamat, f"{r['MERK']} {r['TIPE']}".strip(),
                     r["NO RANGKA"], r["NO MESIN"], m["TAHUN"], r["WARNA"], r["NO HP"], r["EMAIL"]])
    return rows


def rekap_xlsx(records, meta):
    wb = Workbook()
    _sheet(wb, "REKAP DATA BERKAS", REKAP_HEAD, rekap_rows(records, meta))
    return _bytes(wb)


def file_stem(kind: str, n: int, tanggal: str) -> str:
    """DATA_BERKAS_<jumlah>_<DDMMYYYY> (tanggal ISO YYYY-MM-DD masuk)."""
    y, m, d = (tanggal.split("-") + ["", "", ""])[:3] if tanggal else ("", "", "")
    return f"{kind}_{n}" + (f"_{d}{m}{y}" if y else "")


def build_zip(records, warnings, general, meta, bast_no, pdf_stem, page_images: dict, doc_types=None, cells=None, tanggal=""):
    """Return (zip_bytes, daftar_masalah_verifikasi)."""
    refresh_filenames(records)
    tag = bast_tag(bast_no)
    rows = rename_rows(records, meta, pdf_stem)
    checks = verify_all(records, meta, rows, doc_types)
    general = list(general) + [f"VERIFIKASI: {c}" for c in checks]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{file_stem('DATA_BERKAS', len(records), tanggal)}.xlsx", data_berkas(records, warnings, general, cells))
        z.writestr(f"{file_stem('DATA_RENAME', len(records), tanggal)}.xlsx", rename_xlsx(rows))
        z.writestr(f"{file_stem('REKAP_DATA', len(records), tanggal)}.xlsx", rekap_xlsx(records, meta))
        for _, pg, _, new, _ in rows:
            if pg in page_images:
                z.writestr(f"BERKAS_RENAME/{new}", page_images[pg])
    return buf.getvalue(), checks


def pages_zip(page_images: dict, pdf_stem: str) -> bytes:
    """Semua halaman PDF sebagai JPG bernama '<nama pdf>_<halaman>.jpg' (nama lama di DATA RENAME)."""
    prefix = pdf_stem.replace("_", " ")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for pg, data in sorted(page_images.items()):
            z.writestr(f"{prefix}_{pg}.jpg", data)
    return buf.getvalue()
